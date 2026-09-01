# -*- coding: utf-8 -*-
"""깃 이슈에서 WBS 를 만든다.

운영진 양식의 열한 칸 중 여섯은 이미 깃에 있다. 그것을 그대로 옮기고
계산으로 채울 수 있는 것은 계산한다. 사람이 두 번 적지 않게 한다.

    python tools/build_wbs.py                 # 미리보기 (표를 화면에)
    python tools/build_wbs.py --xlsx 경로     # 엑셀로

WBS 1층(작업 영역)과 2층(관문)은 사람이 정한 것이고 여기 상수로 박혀 있다.
3층(세부 작업)만 이슈에서 온다. 이슈 제목이 사람마다 달라도 1·2층이 표를 잡아 준다.
"""
import argparse, datetime, json, subprocess, sys

sys.stdout.reconfigure(encoding='utf-8')

REPOS = ['foothold-project/foothold-lab', 'foothold-project/foothold-go2']
START = datetime.date(2026, 8, 6)          # 프로젝트 시작 (팀 전일제 착수)
END = datetime.date(2026, 12, 11)          # 최종 발표

# 1층. 라벨 이름과 표에 쓸 이름
AREAS = [
    ('A/정책학습',   'A1 정책 학습'),
    ('A/지형씬제작', 'A2 지형·씬 제작'),
    ('A/평가',       'A3 평가'),
    ('A/트윈렌더',   'A4 트윈·렌더'),
    ('B/항법',       'B1 항법'),
    ('B/인지',       'B2 인지'),
    ('C/기록',       'C1 기록'),
    ('C/운영',       'C2 운영'),
]
AREA_NAME = dict(AREAS)
ORDER = {k: i for i, (k, _) in enumerate(AREAS)}

# 2층. 관문과 산출물
GATES = [
    # 9/4 기획발표. GitHub 마일스톤에는 있는데 여기 없어서 그 이슈들이 관문
    #   없는 행이 됐다 (2026-09-01 실측). 웹 타임라인에도 안 나왔다.
    ('기획발표',        datetime.date(2026, 9, 4),   'deliverables/plan/wbs.md'),
    ('A-정책',          datetime.date(2026, 9, 12),  'deliverables/midterm/policy-metrics.md'),
    ('MVP · 중간발표',   datetime.date(2026, 9, 30),  'deliverables/midterm/generalization-report.md'),
    ('NAV · 실기 항법',  datetime.date(2026, 11, 7),  'deliverables/final/nav-report.md'),
    ('FINAL · 최종발표', datetime.date(2026, 12, 11), 'deliverables/final/final-report.md'),
]
GATE_DUE = {g: d for g, d, _ in GATES}
GATE_DOC = {g: p for g, _, p in GATES}

DONE_PCT = {'CLOSED': 100}


def gh(repo):
    out = subprocess.run(
        ['gh', 'issue', 'list', '-R', repo, '--state', 'all', '--limit', '300',
         '--json', 'number,title,state,labels,milestone,assignees,createdAt,closedAt'],
        capture_output=True, text=True, encoding='utf-8')
    if out.returncode != 0:
        raise SystemExit('gh 실패 (%s): %s' % (repo, out.stderr[:200]))
    return json.loads(out.stdout)


def area_of(issue):
    for l in issue['labels']:
        if l['name'] in AREA_NAME:
            return l['name']
    return ''


def d(s):
    return datetime.date.fromisoformat(s[:10]) if s else None


def rows(today):
    span = (END - START).days or 1
    elapsed = round((today - START).days / span * 100)
    out = []
    for repo in REPOS:
        for i in gh(repo):
            a = area_of(i)
            ms = (i['milestone'] or {}).get('title') or ''
            start = d(i['createdAt'])
            due = GATE_DUE.get(ms)
            actual = d(i.get('closedAt'))
            closed = i['state'] == 'CLOSED'
            plan = ''
            if due and start and due > start:
                plan = round((today - start).days / (due - start).days * 100)
            out.append({
                'sort': (ORDER.get(a, 99), ms, i['number']),
                'wbs': AREA_NAME.get(a, '(미분류)'),
                'name': i['title'],
                'gate': ms or '',
                'start': start,
                'due': due,
                'actual': actual,
                'days': (due - start).days + 1 if (due and start) else '',
                'plan': plan,
                'pct': 100 if closed else 0,
                'who': ','.join(x['login'] for x in i['assignees']) or '',
                'doc': GATE_DOC.get(ms, ''),
                'repo': repo.split('/')[-1],
            })
    out.sort(key=lambda r: r['sort'])
    return out, elapsed


HEAD = ['WBS', '작업 이름', '관문', '기간', '시작날짜', '예상 완료 날짜',
        '실제 완료 날짜', '계획율', '완료율', '담당자', '산출물', '저장소']


def to_cells(r):
    f = lambda x: x.isoformat() if x else ''
    p = lambda x: ('%d%%' % x) if x != '' else ''
    return [r['wbs'], r['name'], r['gate'], r['days'], f(r['start']), f(r['due']),
            f(r['actual']), p(r['plan']), p(r['pct']), r['who'], r['doc'], r['repo']]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx')
    ap.add_argument('--today', default=(datetime.datetime.now(datetime.timezone.utc)
                                    + datetime.timedelta(hours=9)).date().isoformat(),
                help='기준일. 기본은 한국 시간 오늘 (UTC 로 두면 자정~09시에 어제가 된다)')
    a = ap.parse_args()
    today = datetime.date.fromisoformat(a.today)

    rs, elapsed = rows(today)
    if not rs:                                    # 조용한 실패 방지
        raise SystemExit('이슈가 하나도 없습니다. WBS 를 만들지 않습니다.')
    unclassified = sum(1 for r in rs if r['wbs'] == '(미분류)')

    print('기간경과율 %d%% (%s ~ %s, 오늘 %s)' % (elapsed, START, END, today))
    print('행 %d개 · 미분류 %d개' % (len(rs), unclassified))
    print()
    print(' | '.join(HEAD))
    for r in rs:
        print(' | '.join(str(c) for c in to_cells(r)))

    if a.xlsx:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise SystemExit('openpyxl 이 없습니다: pip install openpyxl')
        wb = Workbook()
        ws = wb.active
        ws.title = 'WBS'
        ws.append(['기간경과율', '%d%%' % elapsed, '기준일', today.isoformat()])
        ws.append([])
        ws.append(HEAD)
        for c in ws[3]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='DDDDDD')
        for r in rs:
            ws.append(to_cells(r))
        for col, w in zip('ABCDEFGHIJKL', (16, 52, 18, 8, 13, 15, 15, 9, 9, 14, 44, 14)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A4'
        wb.save(a.xlsx)
        print()
        print('저장: %s' % a.xlsx)
        # 되읽어 확인한다
        from openpyxl import load_workbook
        got = load_workbook(a.xlsx)['WBS']
        assert got.max_row == len(rs) + 3, '행 수가 다릅니다'
        print('되읽기 확인: %d행 %d열' % (got.max_row, got.max_column))


if __name__ == '__main__':
    main()
